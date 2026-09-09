"""Upload-time file checks for Resume Management."""

from __future__ import annotations

import io
import zipfile

from app.config import get_settings
from app.resumes.constants import ALLOWED_MIME_TYPES
from app.resumes.errors import FileRejectedError

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_MIME_BY_SUFFIX = {
    ".pdf": PDF,
    ".docx": DOCX,
}


def sniff_mime(filename: str, content_type: str | None) -> str:
    raw = (content_type or "").split(";")[0].strip().lower()
    if raw in ALLOWED_MIME_TYPES:
        return raw
    lower = filename.lower()
    for suffix, mime in _MIME_BY_SUFFIX.items():
        if lower.endswith(suffix):
            return mime
    raise FileRejectedError(
        "Unsupported file type. Upload a PDF or DOCX resume.",
        status_code=415,
        code="UNSUPPORTED_TYPE",
    )


def validate_upload(*, filename: str, mime_type: str, data: bytes) -> None:
    settings = get_settings()
    if mime_type not in ALLOWED_MIME_TYPES:
        raise FileRejectedError(
            "Unsupported file type. Upload a PDF or DOCX resume.",
            status_code=415,
            code="UNSUPPORTED_TYPE",
        )
    if len(data) > settings.resume_max_upload_bytes:
        raise FileRejectedError(
            f"File exceeds {settings.resume_max_upload_bytes} bytes",
            status_code=413,
            code="FILE_TOO_LARGE",
        )
    if not data:
        raise FileRejectedError("File is empty", status_code=400, code="EMPTY_FILE")
    if mime_type == PDF:
        _assert_pdf_ok(data, settings.resume_max_pdf_pages)
    else:
        _assert_docx_ok(data)


def _assert_pdf_ok(data: bytes, max_pages: int) -> None:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise FileRejectedError("PDF support is not installed", status_code=500, code="DEPENDENCY") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise FileRejectedError("PDF is corrupted", status_code=400, code="CORRUPT_FILE") from exc
    if getattr(reader, "is_encrypted", False):
        raise FileRejectedError(
            "Password-protected PDFs are not supported",
            status_code=400,
            code="ENCRYPTED_FILE",
        )
    try:
        page_count = len(reader.pages)
    except Exception:
        page_count = 0
    if page_count > max_pages:
        raise FileRejectedError(
            f"PDF exceeds {max_pages} pages",
            status_code=400,
            code="TOO_MANY_PAGES",
        )


def _assert_docx_ok(data: bytes) -> None:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise FileRejectedError("DOCX is corrupted", status_code=400, code="CORRUPT_FILE")
    try:
        from docx import Document

        Document(io.BytesIO(data))
    except Exception as exc:
        raise FileRejectedError("DOCX is corrupted", status_code=400, code="CORRUPT_FILE") from exc


def extract_text(filename: str, mime_type: str, data: bytes) -> str:
    if mime_type == PDF:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if getattr(reader, "is_encrypted", False):
            raise FileRejectedError(
                "Password-protected PDFs are not supported",
                status_code=400,
                code="ENCRYPTED_FILE",
            )
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    from docx import Document

    document = Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs if p.text).strip()
